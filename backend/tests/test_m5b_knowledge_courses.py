"""M5b việc 9 — kho tri thức hợp nhất (nạp / trích dẫn / phạm vi / lỗ hổng) và khoá học (học → bài học → thi → chứng chỉ)."""
import io
import json

import pytest

from app.models import Agent, Approval, Member
from app.models.learning import CourseEnrollment, KnowledgeGap
from app.services import company_mcp, courses, knowledge_hub
from tests.test_m3_agents import env, hire  # noqa: F401  (fixture)

POLICY = """# Chính sách đổi trả Tiệm Gốm

Khách được đổi trả trong vòng 7 ngày kể từ khi nhận hàng, sản phẩm còn nguyên tem.

Hàng vỡ khi vận chuyển: khách gửi ảnh trong 48 giờ, tiệm gửi bù miễn phí.

Không nhận đổi hàng giảm giá trên 50%."""


@pytest.fixture()
def L(env, monkeypatch):
    monkeypatch.setattr("app.api.learning.get_runtime", lambda: env.gw)
    return env


def add(env, **body):
    r = env.c.post("/api/learning/docs", json=body, headers=env.tok("manager"))
    assert r.status_code == 200, r.text
    return r.json()


def test_chunking_keeps_paragraphs_and_overlap():
    text = "\n\n".join(f"Đoạn {i} " + "chữ " * 60 for i in range(12))
    parts = knowledge_hub.chunk_text(text, size=600, overlap=100)
    assert len(parts) > 2 and all(len(p) <= 800 for p in parts)
    assert parts[1].startswith(parts[0][-100:].lstrip()[:20]) or parts[0][-50:] in parts[1]


def test_ingest_dedupe_search_with_citation_and_gap(L):
    d = add(L, title="Chính sách đổi trả", content=POLICY)
    assert d["chunks"] >= 1 and d["duplicate"] is False and d["scope"] == "Toàn tổ chức"
    assert add(L, title="Bản sao", content=POLICY)["duplicate"] is True
    out = hire(L); ag = L.db.get(Agent, out["agent_id"]); m = L.db.get(Member, ag.member_id)
    ctx = company_mcp.ToolContext(db=L.db, member=m, agent=ag, scopes=["company.knowledge:read"])
    r = company_mcp.TOOLS["company_knowledge_search"].handler(ctx, {"query": "đổi trả trong bao nhiêu ngày"})
    assert r["found"] and r["hits"][0]["document_id"] == d["id"], r
    assert r["hits"][0]["cite"].startswith(f"[KB#{d['id']}§") and "7 ngày" in r["hits"][0]["content"]
    miss = company_mcp.TOOLS["company_knowledge_search"].handler(ctx, {"query": "chính sách bảo hành máy lạnh inverter"})
    assert miss["found"] is False and miss["gap_id"]
    gaps = L.c.get("/api/learning/gaps", headers=L.tok("manager")).json()["items"]
    assert gaps[0]["question"].startswith("chính sách bảo hành") and gaps[0]["who"] == m.name
    company_mcp.TOOLS["company_knowledge_search"].handler(ctx, {"query": "Chính sách bảo hành máy lạnh inverter"})
    assert L.db.query(KnowledgeGap).one().count == 2          # cùng câu hỏi → gộp, đếm lần
    r = L.c.post(f"/api/learning/gaps/{gaps[0]['id']}/resolve", json={"document_id": d["id"]}, headers=L.tok("manager"))
    assert r.status_code == 200 and r.json()["status"] == "resolved"
    assert L.c.post("/api/learning/docs", json={"title": "x", "content": POLICY + " thêm"},
                    headers=L.tok("member")).status_code == 403


def test_scope_agent_private_and_department(L):
    a = hire(L); b = hire(L, name="Minh Khoa", department_id=L.mkt.id)
    ag_a = L.db.get(Agent, a["agent_id"]); ag_b = L.db.get(Agent, b["agent_id"])
    add(L, title="Ghi chú riêng Trang", content="Khách VIP Lan thích men lam, luôn tặng kèm thiệp viết tay.", agent_id=ag_a.id)
    add(L, title="Bảng giá marketing", content="Ngân sách quảng cáo Facebook tháng 11 là 500 đô la Mỹ.",
        department_id=L.mkt.id)
    ma, mb = L.db.get(Member, ag_a.member_id), L.db.get(Member, ag_b.member_id)
    ra = knowledge_hub.search(L.db, L.org.id, "khách VIP Lan men lam", member=ma, agent=ag_a, log_gap=False)
    rb = knowledge_hub.search(L.db, L.org.id, "khách VIP Lan men lam", member=mb, agent=ag_b, log_gap=False)
    assert ra["found"] and not rb["hits"]
    assert not knowledge_hub.search(L.db, L.org.id, "ngân sách quảng cáo Facebook", member=ma, agent=ag_a, log_gap=False)["hits"]
    assert knowledge_hub.search(L.db, L.org.id, "ngân sách quảng cáo Facebook", member=mb, agent=ag_b, log_gap=False)["found"]


def test_upload_docx_and_markdown(L):
    import docx
    buf = io.BytesIO(); dd = docx.Document(); dd.add_heading("Quy trình đóng gói", 1)
    dd.add_paragraph("Bọc ba lớp xốp hơi cho bình gốm cao trên 30 cm."); dd.save(buf)
    r = L.c.post("/api/learning/docs/upload", files={"file": ("dong-goi.docx", buf.getvalue(),
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, headers=L.tok("manager"))
    assert r.status_code == 200, r.text
    assert r.json()["title"] == "dong-goi" and r.json()["chunks"] == 1
    r = L.c.post("/api/learning/docs/upload", files={"file": ("x.exe", b"MZ....", "application/octet-stream")},
                 headers=L.tok("manager"))
    assert r.status_code == 422


def test_resync_marks_course_enrollments_stale(L, monkeypatch):
    pages = {"v": "<html><title>Sổ tay</title><body><nav>menu</nav><p>Giờ mở cửa 8h đến 21h mỗi ngày trong tuần.</p></body></html>"}

    class R:
        def __init__(self, t): self.text, self.content, self.status_code, self.headers = t, t.encode(), 200, {"content-type": "text/html"}

    class C:
        def __init__(self, *a, **k): pass
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def get(self, url): return R(pages["v"])
    monkeypatch.setattr("httpx.Client", C)
    d = add(L, url="https://example.com/so-tay")
    assert d["title"] == "Sổ tay" and d["source_type"] == "url"
    doc = knowledge_hub.get(L.db, L.org.id, d["id"]); assert "menu" not in doc.content
    out = hire(L)
    c = L.c.post("/api/learning/courses", json={"title": "Nhập môn", "document_ids": [d["id"]], "agent_ids": [out["agent_id"]]},
                 headers=L.tok("manager")).json()
    e = L.db.query(CourseEnrollment).one(); e.status, e.module_index = "passed", 1; L.db.commit()
    same = L.c.post(f"/api/learning/docs/{d['id']}/resync", headers=L.tok("manager")).json()
    assert same["changed"] is False
    pages["v"] = pages["v"].replace("21h", "22h")
    ch = L.c.post(f"/api/learning/docs/{d['id']}/resync", headers=L.tok("manager")).json()
    assert ch["changed"] and ch["stale_enrollments"] == 1
    L.db.refresh(e); assert e.status == "stale" and e.module_index == 0 and "thay đổi" in e.note
    assert c["modules"][0]["document_id"] == d["id"]


def test_course_study_lessons_quiz_exam_certificate(L):
    d1 = add(L, title="Chính sách đổi trả", content=POLICY)
    d2 = add(L, title="Giờ làm việc", content="Tiệm mở cửa từ 8 giờ sáng đến 21 giờ tối, nghỉ ngày mùng 1 Tết.")
    out = hire(L); aid = out["agent_id"]
    H = L.tok("manager")
    c = L.c.post("/api/learning/courses", json={"title": "Nhập môn CSKH", "goal": "Trả lời đúng chính sách",
                                               "document_ids": [d1["id"], d2["id"]], "agent_ids": [aid], "pass_score": 80},
                 headers=H)
    assert c.status_code == 200, c.text
    c = c.json(); enr = c["enrollments"][0]
    prompts = []
    quiz = {"questions": [
        {"question": "Khách được đổi trả trong bao lâu?", "options": ["3 ngày", "7 ngày", "14 ngày", "30 ngày"], "answer": 1, "document_id": d1["id"]},
        {"question": "Hàng vỡ phải gửi ảnh trong bao lâu?", "options": ["12 giờ", "24 giờ", "48 giờ", "72 giờ"], "answer": 2, "document_id": d1["id"]},
        {"question": "Tiệm đóng cửa lúc mấy giờ?", "options": ["20h", "21h", "22h", "23h"], "answer": 1, "document_id": d2["id"]},
        {"question": "Sai định dạng", "options": ["a", "b"], "answer": 0}]}
    replies = iter([
        json.dumps(quiz, ensure_ascii=False),
        '{"notes": ["Đổi trả 7 ngày"], "lessons": ["Khi khách xin đổi hàng, hỏi ngày nhận hàng trước để biết còn trong 7 ngày không"]}',
        '{"notes": ["Mở 8h-21h"], "lessons": []}',
        '{"answers": [1, 2, 0]}',          # sai câu 3 → 66.7 < 80
        '{"answers": [1, 2, 1]}',
    ])

    async def chat_stream(agent_id, message, *, session_key=None, timeout_seconds=90):
        prompts.append((session_key, message))
        yield {"kind": "final", "text": next(replies)}
    L.gw.chat_stream = chat_stream
    q = L.c.post(f"/api/learning/courses/{c['id']}/quiz/generate", json={"author_agent_id": aid, "questions": 4}, headers=H)
    assert q.status_code == 200, q.text
    assert len(q.json()["quiz"]) == 3                       # câu sai định dạng bị loại
    s1 = L.c.post(f"/api/learning/enrollments/{enr['id']}/step", headers=H).json()
    assert s1["enrollment"]["module_index"] == 1 and "bài học" in s1["message"]
    assert "7 ngày kể từ khi nhận hàng" in prompts[1][1] and "course-" in prompts[1][0]
    les = L.db.query(Approval).filter(Approval.policy_key.like(f"lesson:a{aid}:%")).one()
    assert json.loads(les.evidence)["source"] == "course"
    s2 = L.c.post(f"/api/learning/enrollments/{enr['id']}/step", headers=H).json()
    assert s2["enrollment"]["status"] == "exam"
    s3 = L.c.post(f"/api/learning/enrollments/{enr['id']}/step", headers=H).json()
    assert s3["enrollment"]["status"] == "failed" and s3["grade"]["wrong"] == [2]
    assert '"answer":' not in prompts[3][1] and "document_id" not in prompts[3][1]   # đề gửi agent không lộ đáp án
    assert L.db.query(KnowledgeGap).filter(KnowledgeGap.source == "exam").count() == 1
    L.c.post(f"/api/learning/enrollments/{enr['id']}/retake", json={}, headers=H)
    s4 = L.c.post(f"/api/learning/enrollments/{enr['id']}/step", headers=H).json()
    assert s4["enrollment"]["status"] == "passed" and s4["enrollment"]["certified_at"] and "ĐẠT" in s4["message"]
    me = L.c.get(f"/api/learning/agents/{aid}", headers=L.tok("member")).json()
    assert me["items"][0]["status"] == "passed" and me["items"][0]["course_title"] == "Nhập môn CSKH"
    assert L.c.post(f"/api/learning/enrollments/{enr['id']}/step", headers=L.tok("member")).status_code == 403


def test_tick_runs_one_step_per_day_at_study_hour(L):
    import asyncio
    from datetime import datetime
    d = add(L, title="Giờ làm việc", content="Tiệm mở cửa từ 8 giờ sáng đến 21 giờ tối, nghỉ ngày mùng 1 Tết.")
    out = hire(L)
    c = courses.create(L.db, type("P", (), {"user_id": None})(), L.org.id,
                       {"title": "Ca tối", "document_ids": [d["id"]], "agent_ids": [out["agent_id"]],
                        "study_hour": 22, "study_tz": "Asia/Ho_Chi_Minh"})
    calls = []

    async def chat_stream(agent_id, message, *, session_key=None, timeout_seconds=90):
        calls.append(session_key); yield {"kind": "final", "text": '{"notes": [], "lessons": []}'}
    L.gw.chat_stream = chat_stream
    assert asyncio.run(courses.tick(L.db, L.gw, now=datetime(2026, 10, 15, 10, 0)))["steps"] == 0   # 17h VN
    assert asyncio.run(courses.tick(L.db, L.gw, now=datetime(2026, 10, 15, 15, 5)))["steps"] == 1   # 22h VN
    assert asyncio.run(courses.tick(L.db, L.gw, now=datetime(2026, 10, 15, 15, 40)))["steps"] == 0  # đã học hôm nay
    assert c["enrollments"] and len(calls) == 1


def test_parse_answers_formats():
    assert courses.parse_answers('{"answers": [1, 2, 0]}', 3) == [1, 2, 0]
    assert courses.parse_answers('{"answers": ["B", "C", "A"]}', 3) == [1, 2, 0]
    assert courses.parse_answers("Đáp án: [1, 2, 0]", 3) == [1, 2, 0]
    assert courses.parse_answers("Câu 1: 1\nCâu 2: 2\nCâu 3: 0", 3) == [1, 2, 0]
    assert courses.parse_answers("Em không biết", 3) == []


def test_overlapping_steps_second_result_is_discarded(L):
    d = add(L, title="Giờ làm việc", content="Tiệm mở cửa từ 8 giờ sáng đến 21 giờ tối, nghỉ ngày mùng 1 Tết.")
    out = hire(L)
    c = L.c.post("/api/learning/courses", json={"title": "Ca tối", "document_ids": [d["id"]], "agent_ids": [out["agent_id"]]},
                 headers=L.tok("manager")).json()
    e = L.db.query(CourseEnrollment).one()

    async def chat_stream(agent_id, message, *, session_key=None, timeout_seconds=90):
        # trong lúc agent đang trả lời, một bước khác đã ghi xong
        e.log = list(e.log or []) + [{"kind": "study"}]; e.module_index = 1; e.status = "exam"; L.db.commit()
        yield {"kind": "final", "text": '{"notes": [], "lessons": []}'}
    L.gw.chat_stream = chat_stream
    r = L.c.post(f"/api/learning/enrollments/{c['enrollments'][0]['id']}/step", headers=L.tok("manager"))
    assert r.status_code == 409, r.text


def test_memory_search_also_returns_library_hits(L):
    add(L, title="Chính sách đổi trả", content=POLICY)
    out = hire(L); ag = L.db.get(Agent, out["agent_id"]); m = L.db.get(Member, ag.member_id)
    ctx = company_mcp.ToolContext(db=L.db, member=m, agent=ag, scopes=["company.knowledge:read"])
    r = company_mcp.TOOLS["company_memory_search"].handler(ctx, {"query": "hàng vỡ khi vận chuyển gửi ảnh"})
    assert r["knowledge"] and "48 giờ" in r["knowledge"][0]["content"] and r["knowledge"][0]["cite"].startswith("[KB#")


def test_core_knowledge_three_layers_compiled_into_agents_md(L):
    a = hire(L)                                   # chức danh: "Nhân viên chăm sóc khách"
    b = hire(L, name="Minh Khoa", role="Nhân viên kho")
    ag_a = L.db.get(Agent, a["agent_id"])
    add(L, title="Chính sách đổi trả", content=POLICY, is_core=True)
    add(L, title="Kịch bản CSKH", content="Luôn chào khách bằng tên và hỏi mã đơn trước khi xử lý.", is_core=True,
        role_keyword="chăm sóc")
    add(L, title="Ghi chú riêng Trang", content="Khách VIP Lan thích men lam, tặng kèm thiệp viết tay.", is_core=True,
        agent_id=ag_a.id)
    add(L, title="Tài liệu tra cứu", content="Bảng mã màu men chi tiết dùng khi khách hỏi sâu về kỹ thuật nung.")
    pa = L.c.get(f"/api/learning/core/{a['agent_id']}", headers=L.tok("member")).json()
    pb = L.c.get(f"/api/learning/core/{b['agent_id']}", headers=L.tok("member")).json()
    assert [d["layer"] for d in pa["documents"]] == ["company", "profession", "personal"]
    assert "### Riêng bạn" in pa["text"] and "### Nghề nghiệp" in pa["text"] and "7 ngày" in pa["text"]
    assert "Bảng mã màu" not in pa["text"]                                    # không cốt lõi → chỉ tra cứu
    assert [d["layer"] for d in pb["documents"]] == ["company"] and "Kịch bản" not in pb["text"]
    r = L.c.post("/api/learning/core/sync", headers=L.tok("manager"))
    assert r.status_code == 200 and r.json()["written"] >= 2, r.text
    agents_md = L.gw.files[(a["runtime_agent_id"], "AGENTS.md")]
    assert "## Kiến thức cốt lõi" in agents_md and "Khách VIP Lan" in agents_md and "hỏi mã đơn" in agents_md
    assert "Khách VIP Lan" not in L.gw.files[(b["runtime_agent_id"], "AGENTS.md")]
    # lớp nghề cũng lọc tra cứu: seat kho không thấy kịch bản CSKH
    mb = L.db.get(Member, L.db.get(Agent, b["agent_id"]).member_id)
    assert not knowledge_hub.search(L.db, L.org.id, "kịch bản chào khách mã đơn", member=mb,
                                    agent=L.db.get(Agent, b["agent_id"]), log_gap=False)["found"]
    # tắt cốt lõi → bỏ khỏi AGENTS.md sau lần sync
    doc_id = next(d["id"] for d in pa["documents"] if d["layer"] == "personal")
    r = L.c.patch(f"/api/learning/docs/{doc_id}", json={"is_core": False}, headers=L.tok("manager"))
    assert r.json()["core_sync"] is True
    assert "Khách VIP Lan" not in L.c.get(f"/api/learning/core/{a['agent_id']}", headers=L.tok("member")).json()["text"]
