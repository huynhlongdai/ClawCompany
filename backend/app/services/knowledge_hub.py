"""M5b việc 9 — kho tri thức hợp nhất: nạp (văn bản / tệp / URL) → cắt đoạn → chỉ mục → tra có trích dẫn.

* Một tài liệu = ``KnowledgeDocument`` (phạm vi: toàn tổ chức, công ty, phòng, dự án hoặc riêng một agent).
  Trùng nội dung (``content_hash``) trong cùng tổ chức + phạm vi → trả lại tài liệu cũ, không nạp lại.
* Cắt đoạn theo đoạn văn (~1.000 ký tự, chồng 150) để trích dẫn trỏ đúng chỗ.
* Tra cứu: điểm lai = độ phủ từ khoá (chính xác với tiếng Việt có dấu) + cosine vector băm 384 chiều sẵn có.
  Agent chỉ thấy tài liệu thuộc phạm vi của mình. Không tìm thấy gì đủ tốt → ghi "lỗ hổng kiến thức".
* Tài liệu từ URL đồng bộ lại được; nội dung đổi → các khoá học dùng tài liệu đó chuyển agent sang "cần ôn".
"""
from __future__ import annotations

import hashlib
import io
import re
import json
from datetime import datetime
from html.parser import HTMLParser

from fastapi import HTTPException
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.models import (Agent, Department, KnowledgeChunk, KnowledgeDocument, KnowledgeVector, Member, Project, Task)
from app.models.learning import Course, CourseEnrollment, KnowledgeGap
from app.services.vector_search import TOKEN_RE, cosine, hash384_embedding, index_chunk_vector

CHUNK = 1000
OVERLAP = 150
MAX_CHARS = 400_000
MAX_FETCH = 3_000_000
GAP_THRESHOLD = 0.34      # điểm tốt nhất dưới mức này = coi như không có tài liệu trả lời được
STOP = {"và", "là", "của", "có", "cho", "các", "những", "một", "được", "không", "với", "trong", "khi", "thì",
        "này", "đó", "the", "a", "an", "of", "to", "is", "in", "and", "or", "for", "gì", "nào", "sao", "bao"}


# ------------------------------------------------------------------ đọc tệp / URL

class _Html(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "nav", "footer", "header", "form"}
    BLOCK = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article", "pre", "td"}

    def __init__(self):
        super().__init__()
        self.out, self.skip, self.title, self._t = [], 0, "", False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        if tag == "title":
            self._t = True
        if tag in self.BLOCK:
            self.out.append("\n")
        if tag in ("h1", "h2", "h3") and not self.skip:
            self.out.append("## ")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        if tag == "title":
            self._t = False
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if self._t:
            self.title += data
        elif not self.skip:
            self.out.append(data)


def html_to_text(html: str) -> tuple[str, str]:
    p = _Html()
    p.feed(html)
    text = re.sub(r"[ \t\r\f\v]+", " ", "".join(p.out))
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    return p.title.strip(), text


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n\n".join((pg.extract_text() or "").strip() for pg in reader.pages).strip()
        if name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = []
            for para in d.paragraphs:
                t = para.text.strip()
                if t:
                    parts.append(("## " + t) if (para.style is not None and str(para.style.name).lower().startswith("heading")) else t)
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n\n".join(parts)
        raw = data.decode("utf-8", errors="replace")
        if name.endswith((".html", ".htm")):
            return html_to_text(raw)[1]
        if name.endswith((".txt", ".md", ".markdown", ".csv", ".json", ".yaml", ".yml")) or "." not in name:
            return raw
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(422, f"Không đọc được tệp {filename}: {exc}") from exc
    raise HTTPException(422, "Chỉ nhận PDF, DOCX, Markdown, TXT, CSV, HTML")


def fetch_url(url: str) -> tuple[str, str]:
    import httpx
    if not re.match(r"^https?://", url or ""):
        raise HTTPException(422, "URL phải bắt đầu bằng http:// hoặc https://")
    try:
        with httpx.Client(follow_redirects=True, timeout=20, headers={"User-Agent": "ClawCompanyKB/1.0"}) as c:
            r = c.get(url)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(502, f"Không tải được URL: {exc}") from exc
    if r.status_code >= 400:
        raise HTTPException(502, f"URL trả lỗi HTTP {r.status_code}")
    if len(r.content) > MAX_FETCH:
        raise HTTPException(422, "Trang quá lớn (tối đa 3 MB)")
    ctype = r.headers.get("content-type", "")
    if "pdf" in ctype:
        return "", extract_text("x.pdf", r.content)
    if "html" in ctype or r.text.lstrip().startswith("<"):
        return html_to_text(r.text)
    return "", r.text


# ------------------------------------------------------------------ cắt đoạn + chỉ mục

def chunk_text(text: str, size: int = CHUNK, overlap: int = OVERLAP) -> list[str]:
    paras = [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]
    out, cur = [], ""
    for p in paras:
        while len(p) > size:                       # đoạn quá dài → cắt cứng, giữ chồng lấn
            if cur:
                out.append(cur); cur = ""
            out.append(p[:size]); p = p[size - overlap:]
        if len(cur) + len(p) + 2 > size and cur:
            out.append(cur)
            tail = cur[-overlap:]
            cur = (tail + "\n\n" + p) if overlap else p
        else:
            cur = (cur + "\n\n" + p) if cur else p
    if cur:
        out.append(cur)
    return out


def _hash(text: str) -> str:
    return hashlib.sha256(re.sub(r"\s+", " ", text or "").strip().encode()).hexdigest()


def reindex(db: Session, doc: KnowledgeDocument) -> int:
    old = [c.id for c in db.query(KnowledgeChunk.id).filter(KnowledgeChunk.document_id == doc.id).all()]
    if old:
        db.query(KnowledgeVector).filter(KnowledgeVector.chunk_id.in_(old)).delete(synchronize_session=False)
        db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == doc.id).delete(synchronize_session=False)
    parts = chunk_text(doc.content)
    from app.services import embeddings
    real = embeddings.embed(parts)
    prov = embeddings.provider_name() if real else "hash384"
    for i, part in enumerate(parts):
        ch = KnowledgeChunk(document_id=doc.id, chunk_index=i, content=part, embedding_provider=prov)
        db.add(ch); db.flush()
        if real:
            db.add(KnowledgeVector(chunk_id=ch.id, embedding_json=json.dumps(real[i])))
        else:
            index_chunk_vector(db, ch)
    doc.indexed = True
    db.add(doc); db.commit()
    return len(parts)


def _check_scope(db: Session, org_id: int, data: dict) -> dict:
    out = {"company_id": data.get("company_id"), "department_id": data.get("department_id"),
           "project_id": data.get("project_id"), "agent_id": data.get("agent_id"),
           "role_keyword": (data.get("role_keyword") or "").strip()[:120]}
    if out["department_id"]:
        d = db.get(Department, out["department_id"])
        if d is None:
            raise HTTPException(404, "Không thấy phòng ban")
        out["company_id"] = out["company_id"] or d.company_id
    if out["project_id"] and db.get(Project, out["project_id"]) is None:
        raise HTTPException(404, "Không thấy dự án")
    if out["agent_id"]:
        ag = db.get(Agent, out["agent_id"])
        m = db.get(Member, ag.member_id) if ag else None
        if m is None or m.organization_id != org_id:
            raise HTTPException(404, "Không thấy agent")
    return out


def ingest(db: Session, org_id: int, *, title: str, content: str, source_type: str = "text",
           source_url: str = "", scope: dict | None = None, is_core: bool = False) -> dict:
    content = (content or "").strip()
    if len(content) < 20:
        raise HTTPException(422, "Tài liệu trống hoặc quá ngắn (cần ít nhất 20 ký tự có chữ)")
    if len(content) > MAX_CHARS:
        content = content[:MAX_CHARS]
    sc = _check_scope(db, org_id, scope or {})
    h = _hash(content)
    q = db.query(KnowledgeDocument).filter(KnowledgeDocument.organization_id == org_id, KnowledgeDocument.content_hash == h)
    for k, v in sc.items():
        q = q.filter(getattr(KnowledgeDocument, k) == v) if v is not None else q.filter(getattr(KnowledgeDocument, k).is_(None))
    old = q.first()
    if old is not None:
        if is_core and not old.is_core:
            old.is_core = True; db.add(old); db.commit()
        return {**public(db, old), "duplicate": True}
    doc = KnowledgeDocument(organization_id=org_id, title=(title or "Tài liệu").strip()[:220], content=content,
                            source_type=source_type, source_url=source_url[:1000], content_hash=h,
                            synced_at=datetime.utcnow(), access_level="restricted", is_core=bool(is_core), **sc)
    db.add(doc); db.commit(); db.refresh(doc)
    reindex(db, doc)
    return {**public(db, doc), "duplicate": False}


def resync(db: Session, org_id: int, doc_id: int) -> dict:
    doc = get(db, org_id, doc_id)
    if not doc.source_url:
        raise HTTPException(409, "Tài liệu này không có nguồn URL để đồng bộ lại")
    title, text = fetch_url(doc.source_url)
    h = _hash(text)
    doc.synced_at = datetime.utcnow()
    if h == doc.content_hash:
        db.add(doc); db.commit()
        return {**public(db, doc), "changed": False, "stale_enrollments": 0}
    doc.content, doc.content_hash = text[:MAX_CHARS], h
    reindex(db, doc)
    stale = mark_courses_stale(db, doc)
    return {**public(db, doc), "changed": True, "stale_enrollments": stale}


def mark_courses_stale(db: Session, doc: KnowledgeDocument) -> int:
    n = 0
    for c in db.query(Course).filter(Course.organization_id == doc.organization_id).all():
        ids = list(c.document_ids or [])
        if doc.id not in ids:
            continue
        idx = ids.index(doc.id)
        for e in db.query(CourseEnrollment).filter(CourseEnrollment.course_id == c.id).all():
            if e.status in ("passed", "failed", "exam") or e.module_index > idx:
                e.status, e.module_index = "stale", idx
                e.note = f"Tài liệu \"{doc.title}\" vừa thay đổi — cần ôn lại từ module {idx + 1}"
                db.add(e); n += 1
    db.commit()
    return n


def get(db: Session, org_id: int, doc_id: int) -> KnowledgeDocument:
    doc = db.get(KnowledgeDocument, doc_id)
    if doc is None or doc.organization_id != org_id:
        raise HTTPException(404, "Không thấy tài liệu")
    return doc


def delete(db: Session, org_id: int, doc_id: int) -> dict:
    doc = get(db, org_id, doc_id)
    old = [c.id for c in db.query(KnowledgeChunk.id).filter(KnowledgeChunk.document_id == doc.id).all()]
    if old:
        db.query(KnowledgeVector).filter(KnowledgeVector.chunk_id.in_(old)).delete(synchronize_session=False)
    db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == doc.id).delete(synchronize_session=False)
    db.query(KnowledgeGap).filter(KnowledgeGap.resolved_document_id == doc.id).update({KnowledgeGap.resolved_document_id: None})
    db.delete(doc); db.commit()
    return {"deleted": doc_id}


def scope_label(db: Session, doc: KnowledgeDocument) -> str:
    if doc.role_keyword and not doc.agent_id:
        return f"Nghề: {doc.role_keyword}"
    if doc.agent_id:
        ag = db.get(Agent, doc.agent_id); m = db.get(Member, ag.member_id) if ag else None
        return f"Riêng {m.name}" if m else "Riêng một agent"
    if doc.project_id:
        p = db.get(Project, doc.project_id); return f"Dự án {p.name}" if p else "Dự án"
    if doc.department_id:
        d = db.get(Department, doc.department_id); return f"Phòng {d.name}" if d else "Phòng ban"
    if doc.company_id:
        return "Công ty"
    return "Toàn tổ chức"


def public(db: Session, doc: KnowledgeDocument) -> dict:
    n = db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == doc.id).count()
    return {"id": doc.id, "title": doc.title, "source_type": doc.source_type, "source_url": doc.source_url or "",
            "chars": len(doc.content or ""), "chunks": n, "scope": scope_label(db, doc),
            "company_id": doc.company_id, "department_id": doc.department_id, "project_id": doc.project_id,
            "agent_id": doc.agent_id, "is_core": bool(doc.is_core), "role_keyword": doc.role_keyword or "",
            "layer": layer_of(doc), "synced_at": doc.synced_at.isoformat() if doc.synced_at else None,
            "created_at": doc.created_at.isoformat() if getattr(doc, "created_at", None) else None}


def list_docs(db: Session, org_id: int, company_id: int | None = None) -> list[dict]:
    from app.core import company_scope
    q = db.query(KnowledgeDocument).filter(KnowledgeDocument.organization_id == org_id)
    q = company_scope.scope_filter(q, KnowledgeDocument.company_id, company_id)   # R1: theo công ty
    rows = q.order_by(KnowledgeDocument.id.desc()).limit(300).all()
    return [public(db, d) for d in rows]


# ------------------------------------------------------------------ tra cứu có trích dẫn

def _tokens(text: str) -> list[str]:
    return [t for t in TOKEN_RE.findall((text or "").lower()) if len(t) > 1 and t not in STOP]


def visible_docs_query(db: Session, org_id: int, member: Member | None, agent: Agent | None):
    q = db.query(KnowledgeDocument).filter(KnowledgeDocument.organization_id == org_id)
    if member is None:      # người: thấy cả tổ chức (đã qua quyền API)
        return q
    from app.services import seat_leases   # R2: ghế thuê thấy thêm ĐÚNG các tài liệu bên cho thuê cấp
    ls = seat_leases.lease_of_member(db, member.id, live_only=True)
    scope = seat_leases.scope_doc_ids(ls) if ls is not None and ls.status == "active" else []
    proj_ids = [r[0] for r in db.query(Task.project_id).filter(Task.assignee_member_id == member.id).distinct().all() if r[0]]
    q = q.filter(or_(KnowledgeDocument.company_id.is_(None), KnowledgeDocument.company_id == member.company_id))
    q = q.filter(or_(KnowledgeDocument.department_id.is_(None), KnowledgeDocument.department_id == member.department_id))
    q = q.filter(or_(KnowledgeDocument.project_id.is_(None), KnowledgeDocument.project_id.in_(proj_ids or [-1])))
    q = q.filter(or_(KnowledgeDocument.agent_id.is_(None), KnowledgeDocument.agent_id == (agent.id if agent else -1)))
    if scope:
        own = [r[0] for r in q.with_entities(KnowledgeDocument.id).all()]
        return db.query(KnowledgeDocument).filter(KnowledgeDocument.id.in_(own + scope))
    return q


def search(db: Session, org_id: int, query: str, *, member: Member | None = None, agent: Agent | None = None,
           limit: int = 6, log_gap: bool = True, task_id: int | None = None) -> dict:
    import json as _json
    query = (query or "").strip()
    qt = set(_tokens(query))
    docs = {d.id: d for d in visible_docs_query(db, org_id, member, agent).all() if role_ok(d, member)}
    if not qt or not docs:
        hits = []
    else:
        from app.services import embeddings
        qemb = hash384_embedding(query)
        prov = embeddings.provider_name()
        real_q = (embeddings.embed([query]) or [None])[0] if prov != "hash384" else None
        rows = (db.query(KnowledgeChunk, KnowledgeVector)
                .outerjoin(KnowledgeVector, KnowledgeVector.chunk_id == KnowledgeChunk.id)
                .filter(KnowledgeChunk.document_id.in_(list(docs))).limit(20000).all())
        scored = []
        for ch, kv in rows:
            ct = set(_tokens(ch.content))
            cover = len(qt & ct) / len(qt)
            title_bonus = 0.1 if qt & set(_tokens(docs[ch.document_id].title)) else 0.0
            try:
                emb = _json.loads(kv.embedding_json) if kv else []
            except Exception:  # noqa: BLE001
                emb = []
            if real_q is not None and ch.embedding_provider == prov and emb:
                # Embedding thật: hiểu cùng ý khác chữ ("bể" ~ "vỡ") → cho vào cả khi không trùng từ nào.
                sem = embeddings.semantic(cosine(real_q, emb))
                s = 0.45 * cover + 0.55 * sem + title_bonus
                if cover > 0 or sem >= 0.35:
                    scored.append((s, ch))
                continue
            if ch.embedding_provider != "hash384" or not emb:
                emb = hash384_embedding(ch.content)
            s = 0.7 * cover + 0.3 * max(0.0, cosine(qemb, emb)) + title_bonus
            if cover > 0:
                scored.append((s, ch))
        scored.sort(key=lambda x: -x[0])
        hits = []
        for s, ch in scored[:limit]:
            d = docs[ch.document_id]
            hits.append({"chunk_id": ch.id, "document_id": d.id, "title": d.title, "chunk": ch.chunk_index, "score": round(s, 3),
                         "source_url": d.source_url or "", "scope": scope_label(db, d),
                         "cite": f"[KB#{d.id}§{ch.chunk_index + 1}: {d.title}]", "content": ch.content})
    best = hits[0]["score"] if hits else 0.0
    gap = None
    if log_gap and member is not None and query and best < GAP_THRESHOLD:
        gap = record_gap(db, org_id, member.id, query, source="search", task_id=task_id)
    return {"query": query, "hits": hits, "best_score": best, "found": best >= GAP_THRESHOLD,
            "gap_id": gap.id if gap else None}


# ------------------------------------------------------------------ lỗ hổng kiến thức

def record_gap(db: Session, org_id: int, member_id: int | None, question: str, *, source: str,
               task_id: int | None = None) -> KnowledgeGap:
    key = hashlib.sha256(" ".join(sorted(set(_tokens(question)))).encode()).hexdigest()[:32]
    g = (db.query(KnowledgeGap).filter(KnowledgeGap.organization_id == org_id, KnowledgeGap.question_key == key,
                                       KnowledgeGap.status == "open").first())
    if g is not None:
        g.count += 1; g.updated_at = datetime.utcnow()
    else:
        g = KnowledgeGap(organization_id=org_id, member_id=member_id, question=question[:2000], question_key=key,
                         source=source, task_id=task_id)
    db.add(g); db.commit(); db.refresh(g)
    return g


def list_gaps(db: Session, org_id: int, status: str = "open") -> list[dict]:
    q = db.query(KnowledgeGap).filter(KnowledgeGap.organization_id == org_id)
    if status != "all":
        q = q.filter(KnowledgeGap.status == status)
    names = {m.id: m.name for m in db.query(Member).filter(Member.organization_id == org_id).all()}
    out = []
    for g in q.order_by(KnowledgeGap.count.desc(), KnowledgeGap.updated_at.desc()).limit(200).all():
        out.append({"id": g.id, "question": g.question, "who": names.get(g.member_id), "source": g.source,
                    "source_vi": {"search": "Tra không ra", "report": "Agent báo thiếu", "exam": "Sai khi thi"}.get(g.source, g.source),
                    "count": g.count, "status": g.status, "task_id": g.task_id,
                    "resolved_document_id": g.resolved_document_id,
                    "updated_at": g.updated_at.isoformat() if g.updated_at else None})
    return out


def resolve_gap(db: Session, org_id: int, gap_id: int, *, document_id: int | None, dismiss: bool = False) -> dict:
    g = db.get(KnowledgeGap, gap_id)
    if g is None or g.organization_id != org_id:
        raise HTTPException(404, "Không thấy lỗ hổng này")
    if document_id:
        get(db, org_id, document_id)
    g.status = "dismissed" if dismiss else "resolved"
    g.resolved_document_id = document_id
    g.resolved_at = datetime.utcnow()
    db.add(g); db.commit()
    return {"id": g.id, "status": g.status}


# ------------------------------------------------------------------ kiến thức cốt lõi 3 lớp → AGENTS.md
# Đo trên VM (eval CSKH 15/10/2026): agent KHÔNG tự gọi company_knowledge_search khi khách hỏi chính sách,
# kể cả khi AGENTS.md dặn — chỉ tra khi được nhắc "dựa theo tài liệu". Nên phần cốt lõi phải nằm sẵn trong
# ngữ cảnh (bài học FastClaw: kho nhỏ thì nhúng thẳng). Ba lớp: riêng agent / nghề (chức danh) / công ty.

CORE_BUDGET = 12_000
LAYER_BUDGET = {"personal": 3_000, "profession": 4_500, "company": 4_500}
LAYER_VI = {"personal": "Riêng bạn", "profession": "Nghề nghiệp", "company": "Công ty"}


def layer_of(doc: KnowledgeDocument) -> str:
    if doc.agent_id:
        return "personal"
    if doc.role_keyword:
        return "profession"
    return "company"


def role_ok(doc: KnowledgeDocument, member: Member | None) -> bool:
    kw = (doc.role_keyword or "").strip().lower()
    return not kw or member is None or kw in (member.role or "").lower()


def core_docs(db: Session, org_id: int, *, company_id: int | None, department_id: int | None, role: str,
              agent_id: int | None) -> list[KnowledgeDocument]:
    q = (db.query(KnowledgeDocument).filter(KnowledgeDocument.organization_id == org_id, KnowledgeDocument.is_core.is_(True),
                                            KnowledgeDocument.project_id.is_(None))
         .filter(or_(KnowledgeDocument.company_id.is_(None), KnowledgeDocument.company_id == company_id))
         .filter(or_(KnowledgeDocument.department_id.is_(None), KnowledgeDocument.department_id == department_id))
         .filter(or_(KnowledgeDocument.agent_id.is_(None), KnowledgeDocument.agent_id == (agent_id or -1))))
    r = (role or "").lower()
    return [d for d in q.order_by(KnowledgeDocument.id).all() if not d.role_keyword or d.role_keyword.lower() in r]


def core_section(db: Session, org_id: int, *, company_id: int | None, department_id: int | None, role: str,
                 agent_id: int | None) -> str:
    docs = core_docs(db, org_id, company_id=company_id, department_id=department_id, role=role, agent_id=agent_id)
    # M2b: hiến chương công ty đứng đầu phần cốt lõi (không bị cắt bởi CORE_BUDGET ở cuối).
    from app.services.company_foundation import charter_block
    charter = charter_block(db, company_id)
    if not docs and not charter:
        return ""
    parts = ["## Kiến thức cốt lõi (luôn áp dụng khi trả lời; chi tiết khác: company_knowledge_search)"]
    if charter:
        parts.append(charter)
    for layer in ("personal", "profession", "company"):
        mine = [d for d in docs if layer_of(d) == layer]
        if not mine:
            continue
        left = LAYER_BUDGET[layer]
        head = LAYER_VI[layer] + (f" — {role}" if layer == "profession" and role else "")
        parts.append(f"### {head}")
        for d in mine:
            body = re.sub(r"\n{3,}", "\n\n", (d.content or "").strip())
            if left <= 200:
                parts.append(f"- {d.title}: đọc bằng company_knowledge_get document_id={d.id}")
                continue
            take = body[:left - 120]
            more = len(body) > len(take)
            parts.append(f"#### {d.title}\n{take}" + (f"\n(… còn tiếp: company_knowledge_get document_id={d.id})" if more else ""))
            left -= len(take) + len(d.title) + 20
    text = "\n\n".join(parts)
    return text[:CORE_BUDGET]


def core_for_member(db: Session, member: Member, agent: Agent | None) -> str:
    core = core_section(db, member.organization_id, company_id=member.company_id, department_id=member.department_id,
                        role=member.role or "", agent_id=agent.id if agent else None)
    from app.services import seat_leases   # R2: ghế thuê → sàn bên cho thuê đứng trước lớp bên thuê
    return seat_leases.layered_core(db, member, core)


def update_doc(db: Session, org_id: int, doc_id: int, data: dict) -> tuple[dict, bool]:
    """Sửa tiêu đề / cờ cốt lõi / lớp nghề. Trả (tài liệu, có ảnh hưởng tới phần cốt lõi không)."""
    doc = get(db, org_id, doc_id)
    before = (doc.is_core, doc.role_keyword or "", doc.title)
    if data.get("title"):
        doc.title = data["title"].strip()[:220]
    if data.get("is_core") is not None:
        doc.is_core = bool(data["is_core"])
    if data.get("role_keyword") is not None:
        doc.role_keyword = data["role_keyword"].strip()[:120]
    db.add(doc); db.commit()
    after = (doc.is_core, doc.role_keyword or "", doc.title)
    return public(db, doc), (before != after and (before[0] or after[0]))


def core_preview(db: Session, org_id: int, agent_id: int) -> dict:
    from app.services import agent_hr
    ag, m = agent_hr.seat(db, org_id, agent_id)
    text = core_for_member(db, m, ag)
    docs = core_docs(db, org_id, company_id=m.company_id, department_id=m.department_id, role=m.role or "", agent_id=ag.id)
    return {"agent_id": ag.id, "chars": len(text), "budget": CORE_BUDGET, "text": text,
            "documents": [{"id": d.id, "title": d.title, "layer": layer_of(d), "layer_vi": LAYER_VI[layer_of(d)]} for d in docs]}


async def sync_core(db: Session, principal, org_id: int, runtime) -> dict:
    """Sinh lại AGENTS.md cho mọi seat đang làm (kèm phần cốt lõi mới)."""
    from app.services import agent_hr
    seats = (db.query(Agent, Member).join(Member, Member.id == Agent.member_id)
             .filter(Member.organization_id == org_id, Agent.lifecycle.notin_(("retired", "runtime_missing"))).all())
    rows = []
    for ag, m in seats:
        try:
            out = await agent_hr.update_hr(db, principal, org_id, runtime, ag.id, {"regenerate_files": ["AGENTS.md"]}, set())
            w = ((out.get("files") or {}).get("AGENTS.md") or {})
            rows.append({"agent_id": ag.id, "name": m.name, "written": bool(w.get("written")), "note": w.get("reason") or ""})
        except Exception as exc:  # noqa: BLE001
            rows.append({"agent_id": ag.id, "name": m.name, "written": False, "note": str(getattr(exc, "detail", exc))[:200]})
    return {"seats": rows, "written": sum(r["written"] for r in rows)}



def reembed_all(db: Session, org_id: int) -> dict:
    """Chỉ mục lại mọi tài liệu chưa dùng embedding hiện hành (sau khi bật embedding thật hoặc đổi model)."""
    from app.services import embeddings
    prov = embeddings.provider_name()
    docs = db.query(KnowledgeDocument).filter(KnowledgeDocument.organization_id == org_id).all()
    done = skipped = 0
    for d in docs:
        provs = {p for (p,) in db.query(KnowledgeChunk.embedding_provider).filter(KnowledgeChunk.document_id == d.id).distinct()}
        if provs == {prov}:
            skipped += 1
            continue
        reindex(db, d)
        done += 1
    return {"provider": prov, "reindexed": done, "skipped": skipped}


def embedding_status(db: Session, org_id: int) -> dict:
    from app.services import embeddings
    prov = embeddings.provider_name()
    q = (db.query(KnowledgeChunk.embedding_provider).join(KnowledgeDocument, KnowledgeDocument.id == KnowledgeChunk.document_id)
         .filter(KnowledgeDocument.organization_id == org_id))
    total = q.count()
    current = q.filter(KnowledgeChunk.embedding_provider == prov).count()
    return {"provider": prov, "semantic": prov != "hash384", "chunks": total, "chunks_current": current}
